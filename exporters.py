from io import BytesIO
from pathlib import Path

from docx import Document
from fpdf import FPDF


FONT_PATH = Path(__file__).parent / "fonts" / "ARIALUNI.TTF"


def export_txt(content: str) -> bytes:
    return content.encode("utf-8")


def export_docx(content: str, document_type: str) -> bytes:
    document = Document()

    document.add_heading(
        document_type,
        level=1,
    )

    for paragraph in content.split("\n"):
        paragraph = paragraph.strip()

        if paragraph:
            document.add_paragraph(paragraph)

    output = BytesIO()
    document.save(output)

    return output.getvalue()


def export_pdf(content: str, document_type: str) -> bytes:
    pdf = FPDF()

    pdf.set_auto_page_break(
        auto=True,
        margin=15,
    )

    pdf.add_page()

    pdf.add_font(
        "ArialUnicode",
        fname=str(FONT_PATH),
    )

    pdf.set_font(
        "ArialUnicode",
        size=16,
    )

    pdf.multi_cell(
        0,
        10,
        document_type,
    )

    pdf.ln(5)

    pdf.set_font(
        "ArialUnicode",
        size=11,
    )

    for paragraph in content.split("\n"):
        paragraph = paragraph.strip()

        if paragraph:
            pdf.multi_cell(
                0,
                7,
                paragraph,
            )

            pdf.ln(2)

    return bytes(pdf.output())